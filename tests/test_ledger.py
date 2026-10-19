"""核心邏輯測試：資料驗證、儲存、帳單解析、分析與 Excel 匯出。"""
from __future__ import annotations

import datetime as dt
from decimal import Decimal

import pytest

from ledger import analytics, excel_export, statement, storage
from ledger.models import Record, ValidationError


def make(n: int = 0, **kw) -> Record:
    base = dict(id=n or 1, date="2026-10-05", category="餐飲", item="午餐", amount=120, kind="支出")
    base.update(kw)
    return Record(**base)


# ---------- 資料驗證 ----------

def test_ym_is_derived_from_date():
    assert make(date="2026-07-31").ym == "2026-07"


def test_amount_must_be_positive():
    with pytest.raises(ValidationError):
        make(amount=0)


def test_future_date_rejected():
    future = dt.date.today() + dt.timedelta(days=1)
    with pytest.raises(ValidationError):
        make(date=future)


def test_bad_kind_rejected():
    with pytest.raises(ValidationError):
        make(kind="轉出")


def test_signed_amount():
    assert make(kind="收入", amount=100).signed_amount == Decimal("100.00")
    assert make(kind="支出", amount=100).signed_amount == Decimal("-100.00")


def test_from_dict_accepts_web_export_fields():
    r = Record.from_dict({"id": 3, "date": "2026-10-01", "cat": "交通", "item": "加油", "amount": 1650, "type": "支出", "pay": "信用卡"})
    assert r.category == "交通" and r.payment == "信用卡"


# ---------- 儲存 ----------

def test_atomic_roundtrip(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    monkeypatch.setattr(storage, "LEDGER_FILE", tmp_path / "ledger.json")
    monkeypatch.setattr(storage, "BACKUP_DIR", tmp_path / "backup")
    storage.save_records([make(1), make(2, amount=500, category="交通")])
    back = storage.load_records()
    assert len(back) == 2
    assert storage.next_id(back) == 3


def test_backup_rotation(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "LEDGER_FILE", tmp_path / "ledger.json")
    monkeypatch.setattr(storage, "BACKUP_DIR", tmp_path / "backup")
    for i in range(8):
        storage.save_records([make(i + 1)])
    assert len(list((tmp_path / "backup").glob("ledger-*.json"))) <= storage.KEEP_BACKUPS


def test_export_csv(tmp_path):
    out = storage.export_csv([make(1)], tmp_path / "x.csv")
    text = out.read_text(encoding="utf-8-sig")
    assert "日期,月份,類別" in text and "午餐" in text


# ---------- 帳單解析 ----------

CFG = {"statement_keywords": {"交通": ["中油"], "餐飲": ["全聯"]}}


def test_parse_common_formats():
    lines = [
        "2026/10/05 全聯福利中心 1,280",
        "10/07 中油加油站 1,650.00",
        "115/10/09 蝦皮購物 980",
        "本期應繳 12,345",
        "這行沒有金額",
    ]
    rows = statement.parse_lines(lines, CFG, default_year=2026)
    assert len(rows) == 3
    assert rows[0].amount == Decimal("1280.00") and rows[0].category == "餐飲"
    assert rows[1].category == "交通"
    assert rows[2].date.year == 2026  # 民國 115 年


def test_summary_lines_skipped():
    assert statement.parse_lines(["最低應繳 1,000", "循環利息 250"], CFG) == []


def test_card_payment_is_not_income_but_refund_is():
    rows = statement.parse_lines(["09/01 09/01 轉 帳 繳 款 TWD -162,434", "08/30 自動扣繳 -5,000",
                                  "09/03 商店退款 -50"], CFG, 2026)
    assert [(row.item, row.kind) for row in rows] == [("商店退款", "收入")]


def test_dedupe_against_existing():
    rows = statement.parse_lines(["10/05 全聯福利中心 120"], CFG, default_year=2026)
    existing = [make(1, date="2026-10-05", item="全聯福利中心", amount=120)]
    assert statement.dedupe(rows, existing) == []


# ---------- 分析 ----------

def recs():
    return [
        Record(id=1, date="2026-08-05", category="薪資", item="薪資", amount=60000, kind="收入"),
        Record(id=2, date="2026-08-06", category="居住", item="房租", amount=18000),
        Record(id=3, date="2026-09-05", category="薪資", item="薪資", amount=60000, kind="收入"),
        Record(id=4, date="2026-09-06", category="居住", item="房租", amount=18000),
        Record(id=5, date="2026-09-10", category="餐飲", item="外食", amount=6000),
    ]


def test_monthly_summary():
    m = analytics.monthly_summary(recs())
    assert [x["ym"] for x in m] == ["2026-08", "2026-09"]
    assert m[1]["expense"] == Decimal("24000.00")
    assert m[1]["net"] == Decimal("36000.00")
    assert round(m[1]["mom"], 1) == pytest.approx(33.3, abs=0.1)


def test_category_breakdown_shares_sum_to_100():
    shares = [c["share"] for c in analytics.category_breakdown(recs())]
    assert round(sum(shares)) == 100


def test_budget_status_flags_overspend():
    st = analytics.budget_status(recs(), {"餐飲": 4000}, "2026-09")
    assert st[0]["over"] is True and st[0]["remaining"] == Decimal("-2000.00")


def test_forecast_non_negative():
    assert analytics.forecast_next_expense(recs()) >= 0


# ---------- Excel ----------

def test_excel_has_summary_and_one_sheet_per_month(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    out = excel_export.export_excel(recs(), tmp_path / "out.xlsx", {"餐飲": 4000})
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["總表", "2026-08", "2026-09"]
    assert wb["總表"]["A1"].value == "日常帳務總表"
    assert wb["2026-09"]["A4"].value == "日期"


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "ROOT", tmp_path)
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(storage, "LEDGER_FILE", tmp_path / "data" / "ledger.json")
    monkeypatch.setattr(storage, "BACKUP_DIR", tmp_path / "data" / "backup")
    monkeypatch.setattr(storage, "CONFIG_FILE", tmp_path / "config.json")
    return tmp_path


@pytest.fixture
def github(sandbox, monkeypatch):
    from ledger.github_sync import GitHubSync, Settings

    client = GitHubSync(Settings(token="test-token-not-a-secret"))
    calls = []
    responses = []

    def request(method, url, **kwargs):
        from types import SimpleNamespace

        calls.append((method, url, kwargs))
        assert responses, "Unexpected network request"
        status, body = responses.pop(0)
        return SimpleNamespace(status_code=status, json=lambda: body)

    monkeypatch.setattr(client.session, "request", request)
    return client, calls, responses


@pytest.mark.parametrize("amount", ["NaN", "Infinity", "-1", "0.001", "bad"])
def test_invalid_amounts(amount):
    with pytest.raises(ValidationError):
        make(amount=amount)


def test_failed_replace_preserves_ledger(sandbox, monkeypatch):
    storage.save_records([make()])
    previous = storage.LEDGER_FILE.read_bytes()

    def fail_replace(*args):
        raise OSError("Simulated disk failure")

    monkeypatch.setattr(storage.os, "replace", fail_replace)
    with pytest.raises(OSError):
        storage.save_records([make(amount=999)])
    assert storage.LEDGER_FILE.read_bytes() == previous
    assert not list(storage.DATA_DIR.glob("*.tmp"))


def test_corrupt_json_not_silently_replaced(sandbox):
    storage.DATA_DIR.mkdir()
    storage.LEDGER_FILE.write_text("{broken", encoding="utf-8")
    with pytest.raises(ValueError):
        storage.load_records()
    assert storage.LEDGER_FILE.read_text() == "{broken"


def test_duplicate_ids_rejected(sandbox):
    with pytest.raises(ValidationError):
        storage.save_records([make(1), make(1, item="晚餐")])
    assert not storage.LEDGER_FILE.exists()


def test_statement_refunds_and_posting_date():
    rows = statement.parse_lines(["10/05 10/07 中油加油站 120", "10/06 商店退款 -50"], CFG, 2026)
    assert rows[0].item == "中油加油站"
    assert rows[1].kind == "收入" and rows[1].amount == Decimal("50.00")


def test_statement_dedupes_same_import():
    rows = statement.parse_lines(["10/05 全聯 120", "10/05 全聯 120"], CFG, 2026)
    assert len(statement.dedupe(rows, [])) == 1


def test_statement_skips_invalid_ocr_rows_when_requested():
    skipped = []
    rows = statement.parse_lines(["10/05 10/06 商店 TWD 00", "10/07 全聯 120", "13/40 商店 50"], CFG, 2026, skipped)
    assert [row.item for row in rows] == ["全聯"] and len(skipped) == 2
    with pytest.raises(ValidationError):
        statement.parse_lines(["10/05 商店 0"], CFG, 2026)


def test_ocr_rows_merge_low_comma_and_cjk_characters():
    from ledger import ocr

    h = 28
    words = [(100, 500, 60, h, "08/26"), (180, 500, 60, h, "08/28"),
             (300, 500, h, h, "全"), (330, 500, h, h, "聯"),
             (900, 500, 70, h, "TWD"), (1100, 500, 14, h, "1"),
             (1114, 514, 8, h, ","), (1136, 500, 42, h, "280"),
             (100, 560, 60, h, "08/27"), (300, 560, h, h, "中"), (1100, 560, 42, h, "650")]
    assert ocr._rows(words) == ["08/26 08/28 全聯 TWD 1,280", "08/27 中 650"]
    rows = statement.parse_lines(ocr._rows(words)[:1], CFG, 2026)
    assert rows[0].amount == Decimal("1280.00") and rows[0].item == "全聯" and rows[0].category == "餐飲"


def test_empty_analytics():
    assert analytics.monthly_summary([]) == []
    assert analytics.category_breakdown([]) == []
    assert analytics.forecast_next_expense([]) == 0


def test_csv_formula_injection_escaped(tmp_path):
    path = storage.export_csv([make(item="=HYPERLINK(1)")], tmp_path / "safe.csv")
    assert "'=HYPERLINK(1)" in path.read_text(encoding="utf-8-sig")


def test_excel_formula_injection_escaped(tmp_path):
    import openpyxl

    path = excel_export.export_excel([make(item="=1+1")], tmp_path / "safe.xlsx")
    workbook = openpyxl.load_workbook(path)
    assert workbook["2026-10"]["C5"].data_type == "s"
    assert workbook["2026-10"]["C5"].value == "=1+1"


def test_github_dry_run_never_connects(github):
    client, calls, responses = github
    client.settings.token = ""
    assert "未連線" in client.push([make()], dry_run=True)
    assert calls == []


def test_github_public_upload_requires_opt_in(github):
    from ledger.github_sync import SyncError

    client, calls, responses = github
    responses.append((200, {"private": False}))
    with pytest.raises(SyncError, match="公開"):
        client.push([make()])
    assert len(calls) == 1


def test_github_push_creates_file_and_remembers_sha(github):
    import base64
    import json

    client, calls, responses = github
    responses.extend([(200, {"private": True}), (404, {}), (201, {"content": {"sha": "new-sha"}})])
    client.push([make()])
    body = calls[-1][2]["json"]
    assert body["branch"] == "main" and "sha" not in body
    payload = json.loads(base64.b64decode(body["content"]))
    assert payload[0]["item"] == "午餐"
    assert client._state()[client._key()] == "new-sha"


def test_github_remote_change_blocks_upload(github):
    from ledger.github_sync import SyncError

    client, calls, responses = github
    client._remember("old-sha")
    responses.extend([(200, {"private": True}), (200, {"type": "file", "sha": "changed-sha"})])
    with pytest.raises(SyncError, match="修改"):
        client.push([make()])
    assert not any(method == "PUT" for method, url, kwargs in calls)


def test_github_update_uses_expected_sha(github):
    client, calls, responses = github
    client._remember("old-sha")
    responses.extend([(200, {"private": True}), (200, {"type": "file", "sha": "old-sha"}),
                      (200, {"content": {"sha": "next-sha"}})])
    client.push([make()])
    assert calls[-1][2]["json"]["sha"] == "old-sha"


@pytest.mark.parametrize("status", [401, 403, 409, 422, 500])
def test_github_errors_do_not_reveal_token(github, status):
    from ledger.github_sync import SyncError

    client, calls, responses = github
    responses.append((status, {"message": client.settings.token}))
    with pytest.raises(SyncError) as caught:
        client.push([make()])
    assert client.settings.token not in str(caught.value)
    assert client.settings.token not in repr(client.settings)


def test_github_pull_backs_up_and_validates(github):
    import base64
    import json

    client, calls, responses = github
    storage.save_records([make(item="original")])
    content = base64.b64encode(json.dumps([make(2).to_dict()]).encode()).decode()
    responses.append((200, {"type": "file", "sha": "remote-sha", "encoding": "base64", "content": content}))
    client.pull(force=True)
    assert storage.load_records()[0].id == 2
    assert len(list(storage.BACKUP_DIR.glob("*.json"))) == 1


def test_github_pull_refuses_to_replace_local_data(github):
    from ledger.github_sync import SyncError

    client, calls, responses = github
    storage.save_records([make()])
    with pytest.raises(SyncError, match="本機已有"):
        client.pull()
    assert calls == []


def test_github_pull_invalid_data_preserves_local(github):
    import base64
    from ledger.github_sync import SyncError

    client, calls, responses = github
    storage.save_records([make()])
    previous = storage.LEDGER_FILE.read_bytes()
    content = base64.b64encode(b"not-json").decode()
    responses.append((200, {"type": "file", "sha": "invalid", "encoding": "base64", "content": content}))
    with pytest.raises(SyncError, match="格式"):
        client.pull(force=True)
    assert storage.LEDGER_FILE.read_bytes() == previous


def test_cli_add_list_export_delete(sandbox, monkeypatch):
    from ledger import cli

    monkeypatch.setattr(cli, "setup_logging", lambda: None)
    assert cli.main(["add", "--date", "2026-10-05", "--category", "餐飲", "--item", "午餐", "--amount", "180"]) == 0
    assert cli.main(["list", "--month", "2026-10", "--keyword", "午餐"]) == 0
    assert len(storage.load_records()) == 1
    path = sandbox / "out.xlsx"
    assert cli.main(["export", "--excel", str(path)]) == 0
    assert path.exists()
    assert cli.main(["delete", "--id", "1", "--yes"]) == 0
    assert storage.load_records() == []


def test_cli_json_import_is_deduplicated(sandbox, monkeypatch):
    import json
    from ledger import cli

    monkeypatch.setattr(cli, "setup_logging", lambda: None)
    path = sandbox / "web.json"
    path.write_text(json.dumps([{"id": 99, "date": "2026-10-01", "cat": "餐飲", "item": "午餐",
                                 "amount": 120, "type": "支出", "pay": "信用卡"}]), encoding="utf-8")
    for attempt in range(2):
        assert cli.main(["import", "--json", str(path), "--yes"]) == 0
    records = storage.load_records()
    assert len(records) == 1 and records[0].id == 1


def test_cli_imports_web_copy_json_object(sandbox, monkeypatch):
    import json
    from ledger import cli

    monkeypatch.setattr(cli, "setup_logging", lambda: None)
    path = sandbox / "web.json"
    path.write_text(json.dumps({"updated": "2026-10-19T00:00:00Z", "categories": ["餐飲"], "budgets": {}, "records": [
        {"id": 1760000000000.123, "date": "2026-10-02", "ym": "2026-10", "cat": "交通", "item": "加油",
         "amount": 1650, "type": "支出", "pay": "信用卡", "note": ""},
        {"id": 3, "date": "2026-10-03", "cat": "居住", "item": "房租", "amount": 18000,
         "type": "支出", "pay": "轉帳", "demo": True},
    ]}), encoding="utf-8")
    assert cli.main(["import", "--json", str(path), "--yes"]) == 0
    records = storage.load_records()
    assert [(row.id, row.item) for row in records] == [(1, "加油")]


def test_cli_rejects_invalid_budget(sandbox, monkeypatch):
    from ledger import cli

    monkeypatch.setattr(cli, "setup_logging", lambda: None)
    assert cli.main(["budget", "--category", "餐飲", "--amount", "NaN"]) == 1
    assert storage.load_config()["budgets"] == {}


def test_categorize_ignores_ocr_spaces():
    config = {"statement_keywords": {"餐飲": ["麥當勞", "超商"], "購物": ["美德耐"]}}
    assert statement.categorize("台 灣 麥 當 勞S 0 K-3 4 6", config) == "餐飲"
    assert statement.categorize("連支*美德耐商場", config) == "購物"
    assert statement.categorize("不明商家", config) == "其他"


def test_config_merges_new_default_keywords(sandbox):
    import json

    storage.CONFIG_FILE.write_text(json.dumps({"statement_keywords": {"餐飲": ["我的店"]}}), encoding="utf-8")
    keywords = storage.load_config()["statement_keywords"]
    assert "我的店" in keywords["餐飲"] and "麥當勞" in keywords["餐飲"] and "中油" in keywords["交通"]


def test_cli_recategorize_only_changes_other(sandbox, monkeypatch):
    from ledger import cli

    monkeypatch.setattr(cli, "setup_logging", lambda: None)
    storage.save_records([make(1, category="其他", item="統一超商-亞醫"),
                          make(2, category="交通", item="麥當勞"),
                          make(3, category="其他", item="不明商家")])
    assert cli.main(["recategorize", "--yes"]) == 0
    assert [row.category for row in storage.load_records()] == ["餐飲", "交通", "其他"]


def test_cli_used_category_cannot_be_removed(sandbox, monkeypatch):
    from ledger import cli

    monkeypatch.setattr(cli, "setup_logging", lambda: None)
    storage.save_records([make()])
    assert cli.main(["category", "remove", "餐飲"]) == 1
    assert "餐飲" in storage.load_config()["categories"]
