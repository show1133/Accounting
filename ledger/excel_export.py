from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import analytics
from .models import Record

HEADER_FILL = PatternFill("solid", fgColor="167D63")
RED = Font(name="Microsoft JhengHei", color="C62828", bold=True)
NUMBER_FORMAT = '#,##0.00;[Red]-#,##0.00'


def write_text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def style_header(sheet, row: int, count: int) -> None:
    for cell in sheet[row][:count]:
        cell.fill = HEADER_FILL
        cell.font = Font(name="Microsoft JhengHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 25


def export_excel(records: list[Record], path: str | Path, budgets: dict | None = None,
                 months=()) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "總表"
    summary.append(["日常帳務總表"])
    summary["A1"].font = Font(name="Microsoft JhengHei", size=18, bold=True)
    summary.append([])
    summary.append(["月份", "收入", "支出", "結餘", "儲蓄率", "支出環比"])
    style_header(summary, 3, 6)
    monthly = analytics.monthly_summary(records, months)
    for row in monthly:
        summary.append([row["ym"], float(row["income"]), float(row["expense"]),
                        float(row["net"]), row["savings_rate"] / 100,
                        row["mom"] / 100 if row["mom"] is not None else None])
    total_income = sum((row["income"] for row in monthly), Decimal("0"))
    total_expense = sum((row["expense"] for row in monthly), Decimal("0"))
    summary.append(["合計", float(total_income), float(total_expense), float(total_income - total_expense),
                    float((total_income - total_expense) / total_income) if total_income else 0])
    for row in summary.iter_rows(min_row=4, max_row=summary.max_row):
        for cell in row[1:4]:
            cell.number_format = NUMBER_FORMAT
            if cell.value is not None and cell.value < 0:
                cell.font = RED
        for cell in row[4:6]:
            cell.number_format = "0.0%;[Red]-0.0%"
    if monthly:
        chart = BarChart()
        chart.title = "每月收支"
        chart.y_axis.title = "新台幣"
        chart.add_data(Reference(summary, min_col=2, max_col=3, min_row=3, max_row=3 + len(monthly)), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=1, min_row=4, max_row=3 + len(monthly)))
        chart.width, chart.height = 21, 10
        summary.add_chart(chart, "H3")
    summary.append([])
    summary.append(["全期支出類別", "金額", "占比"])
    style_header(summary, summary.max_row, 3)
    for row in analytics.category_breakdown(records):
        summary.append([row["category"], float(row["amount"]), row["share"] / 100])
        write_text(summary.cell(summary.max_row, 1), row["category"])
        summary.cell(summary.max_row, 2).number_format = NUMBER_FORMAT
        summary.cell(summary.max_row, 3).number_format = "0.0%"
    if monthly and budgets:
        month = monthly[-1]["ym"]
        summary.append([])
        summary.append([f"{month} 預算", "預算", "實支", "剩餘", "使用率"])
        style_header(summary, summary.max_row, 5)
        for row in analytics.budget_status(records, budgets, month):
            summary.append([row["category"], float(row["budget"]), float(row["actual"]),
                            float(row["remaining"]), row["ratio"] / 100 if row["ratio"] is not None else None])
            write_text(summary.cell(summary.max_row, 1), row["category"])
            for cell in summary[summary.max_row][1:4]:
                cell.number_format = NUMBER_FORMAT
                if row["over"]:
                    cell.font = RED
            summary.cell(summary.max_row, 5).number_format = "0.0%"
    summary.freeze_panes = "B4"
    for column in range(1, 7):
        summary.column_dimensions[get_column_letter(column)].width = 20
    for month in monthly:
        sheet = workbook.create_sheet(month["ym"])
        sheet.append([f'{month["ym"]} 帳務明細'])
        sheet["A1"].font = Font(name="Microsoft JhengHei", size=16, bold=True)
        sheet.append(["收入", float(month["income"]), "支出", float(month["expense"]), "結餘", float(month["net"])])
        for column in (2, 4, 6):
            sheet.cell(2, column).number_format = NUMBER_FORMAT
        sheet.append([])
        sheet.append(["日期", "類別", "項目", "收支", "金額", "支付方式", "備註"])
        style_header(sheet, 4, 7)
        for record in sorted((row for row in records if row.ym == month["ym"]), key=lambda row: (row.date, row.id)):
            sheet.append([record.date, record.category, record.item, record.kind,
                          float(record.amount), record.payment, record.note])
            for column, text in ((2, record.category), (3, record.item), (6, record.payment), (7, record.note)):
                write_text(sheet.cell(sheet.max_row, column), text)
            sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm-dd"
            amount = sheet.cell(sheet.max_row, 5)
            amount.number_format = NUMBER_FORMAT
            if record.kind == "支出" and record.amount >= 10000:
                amount.font = RED
        sheet.auto_filter.ref = f"A4:G{max(4, sheet.max_row)}"
        sheet.freeze_panes = "A5"
        sheet.append(["月度小計", "", "", "收入", float(month["income"])])
        sheet.append(["", "", "", "支出", float(month["expense"])])
        sheet.append(["", "", "", "結餘", float(month["net"])])
        for row in range(sheet.max_row - 2, sheet.max_row + 1):
            sheet.cell(row, 5).number_format = NUMBER_FORMAT
            if sheet.cell(row, 5).value < 0:
                sheet.cell(row, 5).font = RED
        for column, width in enumerate((15, 16, 32, 10, 18, 16, 40), start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.save(path)
    return path